#!/usr/bin/env python3
"""
Script biên dịch file Báo cáo Word tổng quát (Generic Document Generator).
Không hardcode nội dung báo cáo trong code: Script sẽ đọc file '{lab_name}/docxs.json',
tự động định dạng trang bìa, thông tin sinh viên, các đề mục, bảng biểu và chèn hình ảnh
minh chứng để tạo ra file Word hoàn chỉnh lưu vào thư mục 'bin/docs/'.

Cách sử dụng:
1. Chạy với cấu hình mặc định (lab02 -> xuất vào bin/docs/):
   python gendocs.py

2. Chạy với bài lab hoặc đường dẫn file xuất tùy chỉnh:
   python gendocs.py --lab lab02 --out bin/docs/BaoCao_Lab02.docx
"""

import argparse
import json
import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

# ==================== CẤU HÌNH MẶC ĐỊNH ====================
CONFIG = {
    "lab": "lab02",                                 # Tên bài lab (thư mục chứa docxs.json)
    "docxs_filename": "docxs.json",                 # Tên file cấu hình báo cáo
    "output_dir": "bin/docs",                       # Thư mục xuất file báo cáo mặc định
    "output_file": None                             # None = lấy tên file từ docxs.json và lưu vào bin/docs/
}
# ==========================================================

def parse_arguments():
    parser = argparse.ArgumentParser(description="Tự động biên dịch file Báo cáo thực hành Word từ file JSON vào bin/docs/.")
    parser.add_argument("--lab", "-l", default=CONFIG["lab"], help="Tên bài lab (thư mục chứa docxs.json, vd: lab02)")
    parser.add_argument("--out", "-o", default=CONFIG["output_file"], help="Đường dẫn file Word xuất ra (ghi đè cấu hình)")
    parser.add_argument("--config", "-c", default=CONFIG["docxs_filename"], help="Tên file cấu hình docx (mặc định docxs.json)")
    return parser.parse_args()

def locate_docxs_json(lab_name, filename="docxs.json"):
    """Tìm file docxs.json trong các vị trí tương đối phổ biến"""
    possible_paths = [
        os.path.join(lab_name, filename),
        os.path.join("src", "scripts", lab_name, filename),
        os.path.join(os.path.dirname(__file__), lab_name, filename),
        os.path.join(os.path.dirname(__file__), "..", "scripts", lab_name, filename)
    ]
    for p in possible_paths:
        if os.path.exists(p):
            return os.path.abspath(p)
    return None

def resolve_image_path(img_path, lab_name):
    """Tìm đường dẫn ảnh thực tế, tự động tìm trong bin/imgs/{lab_name}/ nếu cần"""
    if os.path.exists(img_path):
        return img_path
    
    bname = os.path.basename(img_path)
    search_paths = [
        os.path.join("bin", "imgs", lab_name, bname),
        os.path.join("docs", "generated_media", bname),
        os.path.join(os.path.dirname(__file__), "..", "..", "bin", "imgs", lab_name, bname)
    ]
    for p in search_paths:
        if os.path.exists(p):
            return os.path.abspath(p)
    return img_path

def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 32, 96)
    return p

def add_subtitle(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = RGBColor(31, 78, 121)
    return p

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 32, 96)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(31, 78, 121)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    return p

def add_body_paragraph(doc, text, bold_prefix=None, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    r.font.italic = italic
    return p

def add_bullet_item(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    return p

def add_code_block(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(40, 40, 40)
    pPr = p._element.get_or_add_pPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F2F2F2"/>')
    pPr.append(shd)
    return p

def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)
    run.font.italic = True
    run.font.bold = True
    run.font.color.rgb = RGBColor(80, 80, 80)
    return p

def add_image_item(doc, raw_img_path, caption, width_inches=6.0, lab_name="lab02"):
    actual_path = resolve_image_path(raw_img_path, lab_name)
    if os.path.exists(actual_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(actual_path, width=Inches(width_inches))
        add_caption(doc, caption)
    else:
        print(f"      [!] Cảnh báo: Không tìm thấy ảnh '{raw_img_path}' (hoặc '{actual_path}'), bỏ qua.")

def add_table_item(doc, headers, rows):
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.name = "Times New Roman"
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="D9E1F2"/>')
        cell._tc.get_or_add_tcPr().append(shd)

    for r_idx, r_data in enumerate(rows, start=1):
        for c_idx, val in enumerate(r_data):
            cell = table.cell(r_idx, c_idx)
            cell.text = val
            cell.paragraphs[0].runs[0].font.name = "Times New Roman"

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def main():
    args = parse_arguments()
    lab_name = args.lab
    config_file = args.config

    print("==========================================================")
    print("      GENERIC DOCS GENERATOR - BIÊN DỊCH BÁO CÁO VÀO BIN/DOCS ")
    print("==========================================================")
    print(f"[*] Bài Lab       : {lab_name}")

    json_path = locate_docxs_json(lab_name, config_file)
    if not json_path:
        print(f"[!] LỖI: Không tìm thấy file '{lab_name}/{config_file}'!")
        sys.exit(1)

    print(f"[*] File cấu hình : {json_path}")

    with open(json_path, "r", encoding="utf-8") as f:
        try:
            cfg = json.load(f)
        except Exception as e:
            print(f"[!] Lỗi khi đọc file JSON: {e}")
            sys.exit(1)

    metadata = cfg.get("metadata", {})
    
    # Xác định đường dẫn file Word xuất ra (mặc định trong bin/docs/)
    if args.out:
        output_path = args.out
    else:
        raw_output = metadata.get("default_output", f"{lab_name}_BaoCao.docx")
        filename = os.path.basename(raw_output)
        output_path = os.path.join(CONFIG["output_dir"], filename)

    print(f"[*] File xuất ra  : {output_path}")
    print("----------------------------------------------------------")

    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # 1. Trang bìa / Tiêu đề
    if "title" in metadata:
        add_title(doc, metadata["title"])
    if "subtitle" in metadata:
        add_subtitle(doc, metadata["subtitle"])

    # Thông tin sinh viên
    student = metadata.get("student", {})
    if student:
        p_info = doc.add_paragraph()
        p_info.paragraph_format.space_after = Pt(8)
        p_info.paragraph_format.line_spacing = 1.2
        items = [
            ("Họ và tên sinh viên: ", student.get("name", "")),
            ("Mã sinh viên: ", student.get("id", "")),
            ("Lớp: ", student.get("class", "")),
            ("Mã lớp học phần: ", student.get("course_code", "")),
            ("Nhóm thực hành: ", student.get("group", "")),
            ("Mạng LAN ảo: ", student.get("network", ""))
        ]
        for label, val in items:
            if val:
                r_lbl = p_info.add_run(label)
                r_lbl.font.name = "Times New Roman"
                r_lbl.font.size = Pt(11)
                r_lbl.font.bold = True

                r_val = p_info.add_run(val + "\n")
                r_val.font.name = "Times New Roman"
                r_val.font.size = Pt(11)

    # 2. Xử lý danh sách thành phần trong nội dung
    content_items = cfg.get("content", [])
    print(f"[*] Đang xử lý {len(content_items)} thành phần nội dung...")

    for item in content_items:
        c_type = item.get("type", "paragraph")
        
        if c_type == "heading_1":
            add_heading_1(doc, item.get("text", ""))
        elif c_type == "heading_2":
            add_heading_2(doc, item.get("text", ""))
        elif c_type == "bullet":
            add_bullet_item(doc, item.get("text", ""))
        elif c_type == "code":
            add_code_block(doc, item.get("text", ""))
        elif c_type == "image":
            src = item.get("src", "")
            caption = item.get("caption", "")
            width = item.get("width", 6.0)
            add_image_item(doc, src, caption, width, lab_name)
        elif c_type == "table":
            headers = item.get("headers", [])
            rows = item.get("rows", [])
            add_table_item(doc, headers, rows)
        else:
            text = item.get("text", "")
            bold_prefix = item.get("bold_prefix", None)
            italic = item.get("italic", False)
            add_body_paragraph(doc, text, bold_prefix, italic)

    # 3. Lưu file docx
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print("\n==========================================================")
    print(f"[+] THÀNH CÔNG! Đã tạo file báo cáo tại: {output_path}")
    print("==========================================================")

if __name__ == "__main__":
    main()
